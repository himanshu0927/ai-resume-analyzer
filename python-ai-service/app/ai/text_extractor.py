import os
import re
from typing import Tuple

def extract_text_from_file(file_path: str, filename: str) -> Tuple[str, str]:
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == ".pdf":
        text = _extract_from_pdf(file_path)
        file_type = "PDF"
    elif ext in [".docx", ".doc"]:
        text = _extract_from_docx(file_path)
        file_type = "DOCX"
    else:
        raise ValueError(f"Unsupported file format: {ext}. Only PDF and DOCX documents are supported.")
    
    clean_text = _clean_text(text)
    if not clean_text or len(clean_text.strip()) < 10:
        raise ValueError("Could not extract readable text from document. File may be empty or an unreadable scanned image.")
        
    return clean_text, file_type

def _extract_from_pdf(file_path: str) -> str:
    text_content = []
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t: text_content.append(t)
    except Exception:
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            for page in reader.pages:
                t = page.extract_text()
                if t: text_content.append(t)
        except Exception as e:
            raise ValueError(f"Failed to read PDF document: {str(e)}")
            
    return "\n".join(text_content)

def _extract_from_docx(file_path: str) -> str:
    try:
        import docx
        doc = docx.Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip(): full_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text: full_text.append(" | ".join(row_text))
        return "\n".join(full_text)
    except Exception as e:
        raise ValueError(f"Failed to read DOCX document: {str(e)}")

def _clean_text(text: str) -> str:
    if not text: return ""
    text = re.sub(r'[\r\t]', ' ', text)
    text = re.sub(r' +', ' ', text)
    text = re.sub(r'\n\s*\n+', '\n', text)
    return text.strip()
